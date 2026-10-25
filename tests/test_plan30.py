"""The 30-day plan: theoretical curve, real ledger, daily review, Excel export."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
import unittest
from dataclasses import replace
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from unittest import mock

from tmbot.analysis import plan30 as plan_module
from tmbot.analysis.plan30 import DayLog, Plan30
from tmbot.models import Direction, Fill
from tests.helpers import RULES, candles
from tests.test_supervisor import build_supervisor

START = date(2026, 10, 26)
UTC = timezone.utc


def fill(deal="d1", *, day=0, hour=10, profit=None, entry=3400.0, price=3401.0,
         size=1.0, commission=None, swap=None, direction=Direction.BUY, risk=2.0):
    ts = datetime(START.year, START.month, START.day, hour, tzinfo=UTC) + timedelta(days=day)
    return Fill(deal_id=deal, epic="GOLD", ts=ts, stage="TP1", size=size, price=price,
                entry_price=entry, direction=direction, initial_risk=risk, fraction=1.0,
                profit=profit, commission=commission, swap=swap)


def plan(**kw):
    return Plan30(start_date=START, **kw)


def value_100(epic):
    return 100.0


class CurveTests(unittest.TestCase):
    def test_day_one_and_the_compounding_follow_the_spec(self):
        curve = plan_module.theoretical_curve(plan())
        self.assertEqual(curve[0], (100.0, 25.0, 125.0))
        self.assertEqual(curve[1], (125.0, 31.25, 156.25))
        self.assertEqual(curve[2], (156.25, 39.06, 195.31))
        self.assertEqual(len(curve), 30)

    def test_the_scenario_is_enormous_which_is_why_it_is_only_a_scenario(self):
        self.assertGreater(plan_module.theoretical_curve(plan())[-1][2], 80_000)

    def test_rounding_matches_a_spreadsheet_half_cents(self):
        self.assertEqual(plan_module.money(61.035), 61.04)
        self.assertEqual(plan_module.money(244.14 * 0.25), 61.04)
        self.assertEqual(plan_module.money(-0.005), -0.01)
        self.assertEqual(plan_module.money(2.675), 2.68)

    def test_money_is_always_two_decimals(self):
        for start, profit, end in plan_module.theoretical_curve(plan(capital=333.33)):
            for value in (start, profit, end):
                self.assertEqual(value, round(value, 2))


class LedgerTests(unittest.TestCase):
    def test_actual_and_theoretical_are_kept_apart(self):
        p = plan()
        p.day(1).net, p.day(2).net = 10.0, -5.5
        table = plan_module.rows(p)
        self.assertEqual((table[0].actual_start, table[0].theoretical_start), (100.0, 100.0))
        self.assertEqual((table[0].actual_end, table[0].theoretical_end), (110.0, 125.0))
        self.assertEqual((table[1].actual_start, table[1].theoretical_start), (110.0, 125.0))
        self.assertEqual(table[1].actual_end, 104.5)
        self.assertEqual(plan_module.actual_balance(p), 104.5)

    def test_future_days_stay_blank_instead_of_being_filled_in(self):
        p = plan()
        p.day(1).net = 3.0
        table = plan_module.rows(p)
        self.assertEqual(table[1].actual_start, 103.0)
        self.assertIsNone(table[1].actual_end)
        self.assertIsNone(table[1].achieved)
        self.assertIsNone(table[5].actual_start)

    def test_a_day_with_no_known_result_blocks_the_chain(self):
        p = plan()
        p.day(2).net = 4.0          # day 1 unknown
        self.assertIsNone(plan_module.rows(p)[1].actual_start)
        self.assertEqual(plan_module.actual_balance(p), 100.0)

    def test_target_achieved_means_the_25_percent_of_the_actual_balance(self):
        p = plan()
        p.day(1).net, p.day(2).net = 25.0, 20.0
        table = plan_module.rows(p)
        self.assertTrue(table[0].achieved)
        self.assertFalse(table[1].achieved)     # 20 < 25% of 125

    def test_a_losing_day_lowers_the_next_start(self):
        p = plan()
        p.day(1).net = -8.0
        self.assertEqual(plan_module.rows(p)[1].actual_start, 92.0)


class DayTests(unittest.TestCase):
    def evaluate(self, fills, start=100.0, value=value_100, p=None):
        return plan_module.evaluate_day(fills, p or plan(), 1, start, UTC, value)

    def test_a_day_with_no_fills_is_a_no_trade_day_at_zero(self):
        log = self.evaluate([])
        self.assertEqual((log.net, log.trades, log.wins, log.losses), (0.0, 0, 0, 0))

    def test_booked_money_is_used_with_costs(self):
        log = self.evaluate([fill(profit=2.0, commission=-0.1, swap=-0.05)])
        self.assertEqual(log.net, 1.85)
        self.assertEqual(log.costs, -0.15)
        self.assertFalse(log.estimated)
        self.assertEqual((log.trades, log.wins, log.losses), (1, 1, 0))

    def test_missing_money_is_estimated_from_the_price_move(self):
        log = self.evaluate([fill(price=3401.0, entry=3400.0, size=0.02)])
        self.assertEqual(log.net, 2.0)          # 1.0 move x 0.02 lots x 100
        self.assertTrue(log.estimated)

    def test_a_short_estimate_has_the_right_sign(self):
        log = self.evaluate([fill(direction=Direction.SELL, price=3399.0, entry=3400.0,
                                  size=0.01)])
        self.assertEqual(log.net, 1.0)

    def test_money_that_cannot_be_worked_out_is_not_invented(self):
        self.assertIsNone(self.evaluate([fill(profit=None)], value=lambda epic: None))

    def test_only_fills_of_that_local_day_count(self):
        log = self.evaluate([fill(day=0, profit=1.0), fill("d2", day=1, profit=50.0)])
        self.assertEqual(log.net, 1.0)

    def test_wins_losses_and_trades_count_deals_not_fills(self):
        log = self.evaluate([fill("a", profit=1.0), fill("a", hour=11, profit=0.5),
                             fill("b", hour=12, profit=-0.8)])
        self.assertEqual((log.trades, log.wins, log.losses), (2, 1, 1))
        self.assertEqual(log.net, 0.7)

    def test_the_drawdown_is_the_worst_fall_of_the_balance_that_day(self):
        log = self.evaluate([fill("a", hour=9, profit=2.0), fill("b", hour=10, profit=-1.0),
                             fill("c", hour=11, profit=-1.0)])
        self.assertAlmostEqual(log.max_dd, round(2.0 / 102.0 * 100, 2))

    def test_the_daily_loss_limit_is_checked_against_the_days_start(self):
        log = self.evaluate([fill("a", profit=-1.0), fill("b", hour=11, profit=-1.0),
                             fill("c", hour=12, profit=-1.5)])
        self.assertTrue(log.loss_limit_hit)
        self.assertIn(plan_module.RULE_LOSS_LIMIT, log.broken)

    def test_exactly_at_the_limit_stops_the_day_but_is_not_a_breach(self):
        log = self.evaluate([fill("a", profit=-1.0), fill("b", hour=11, profit=-1.0),
                             fill("c", hour=12, profit=-1.0)])
        self.assertTrue(log.loss_limit_hit)
        self.assertNotIn(plan_module.RULE_LOSS_LIMIT, log.broken)

    def test_one_trade_losing_more_than_the_risk_limit_is_flagged(self):
        log = self.evaluate([fill(profit=-1.6)])
        self.assertIn(plan_module.RULE_TRADE_RISK, log.broken)
        self.assertNotIn(plan_module.RULE_LOSS_LIMIT, log.broken)

    def test_the_limits_follow_the_current_balance(self):
        log = self.evaluate([fill(profit=-2.5)], start=200.0)   # 1.25% of 200
        self.assertIn(plan_module.RULE_TRADE_RISK, log.broken)
        self.assertEqual(plan_module.daily_limit(plan(), 200.0), 6.0)


class DecisionTests(unittest.TestCase):
    def test_continue_reduce_pause(self):
        p = plan()
        p.day(1).net = 1.0
        self.assertEqual(plan_module.decide(p, 1), plan_module.CONTINUE)
        p.day(2).broken = [plan_module.RULE_TRADE_RISK]
        self.assertEqual(plan_module.decide(p, 2), plan_module.REDUCE)
        p.day(3).loss_limit_hit = True
        self.assertEqual(plan_module.decide(p, 3), plan_module.PAUSE)

    def test_two_losing_days_in_a_row_reduce_risk(self):
        p = plan()
        p.day(1).net, p.day(2).net = -1.0, -0.5
        self.assertEqual(plan_module.decide(p, 2), plan_module.REDUCE)
        p.day(2).net = 0.5
        self.assertEqual(plan_module.decide(p, 2), plan_module.CONTINUE)


class SummaryTests(unittest.TestCase):
    def test_the_final_figures(self):
        p = plan()
        for number, (net, wins, losses, dd, costs, final) in enumerate([
            (5.0, 2, 1, 1.5, -0.2, True), (-2.0, 0, 1, 2.0, -0.1, True),
            (0.0, 0, 0, None, 0.0, True),
        ], start=1):
            log = p.day(number)
            log.net, log.wins, log.losses, log.max_dd, log.costs = net, wins, losses, dd, costs
            log.trades = wins + losses
            log.final = final
            if number == 2:
                log.broken, log.loss_limit_hit = [plan_module.RULE_TRADE_RISK], False
        s = plan_module.summary(p)
        self.assertEqual((s["initial"], s["final"], s["net"]), (100.0, 103.0, 3.0))
        self.assertEqual(s["return_pct"], 3.0)
        self.assertEqual((s["wins"], s["losses"]), (2, 2))
        self.assertEqual(s["win_rate"], 50.0)
        self.assertEqual(s["max_drawdown"], 2.0)
        self.assertEqual(s["costs"], -0.3)
        self.assertEqual((s["days_reviewed"], s["days_compliant"]), (3, 2))
        self.assertEqual(s["no_trade_days"], 1)
        self.assertEqual(s["difference"], round(s["theoretical_final"] - 103.0, 2))

    def test_a_plan_with_no_days_has_no_made_up_rates(self):
        s = plan_module.summary(plan())
        self.assertIsNone(s["win_rate"])
        self.assertIsNone(s["max_drawdown"])
        self.assertEqual(s["net"], 0.0)


class ServiceTests(unittest.TestCase):
    def setUp(self):
        self.supervisor, self.broker, self.store, self.notifier = build_supervisor()
        self.broker.set_rules(replace(RULES, epic="GOLD", value_per_point=100.0))
        self.service = self.supervisor.plan30
        self.now = datetime(2026, 10, 26, 12, tzinfo=UTC)
        patcher = mock.patch("tmbot.manage.plan30_service.utcnow", lambda: self.now)
        patcher.start()
        self.addCleanup(patcher.stop)

    def at(self, day, hour=12):
        self.now = datetime(2026, 10, 26, hour, tzinfo=UTC) + timedelta(days=day)
        self.service._next_step = 0.0

    def started(self, **kw):
        p = plan(**kw)
        p.save(self.store)
        return p

    def messages(self):
        return [m for _, m in self.notifier.messages]

    # ----------------------------------------------------------- commands
    def test_nothing_runs_until_you_start_it(self):
        self.assertIn("No 30-day plan", self.service.command(""))
        self.service.step()
        self.assertEqual(self.messages(), [])

    def test_start_uses_100_by_default_and_a_given_capital(self):
        self.assertIn("capital 100.00 USD", self.service.command("start"))
        loaded = Plan30.load(self.store)
        self.assertEqual((loaded.capital, loaded.start_date), (100.0, START))
        self.assertIn("preserved", self.service.command("start 500"))
        self.assertEqual(Plan30.load(self.store).capital, 100.0)

    def test_start_rejects_nonsense(self):
        self.assertIn("/plan30", self.service.command("start abc"))
        self.assertIn("/plan30", self.service.command("start -5"))
        self.assertIsNone(Plan30.load(self.store))

    def test_status_says_the_scenario_is_not_an_expectation(self):
        self.started()
        text = self.service.command("")
        self.assertIn("Day 1 of 30", text)
        self.assertIn("Actual at start of day 100.00 USD", text)
        self.assertIn("Theoretical end of day 125.00", text)
        self.assertIn("no trades yet", text)
        self.assertIn("Daily loss limit 3.00 USD (0% used)", text)
        self.assertIn("not an expectation", text)

    def test_status_shows_todays_result_from_the_fills(self):
        self.started()
        self.store.record_fill(fill(profit=1.5))
        text = self.service.command("")
        self.assertIn("net +1.50 (1 trade(s))", text)

    def test_result_records_today_and_overwrites_while_the_day_is_open(self):
        self.started()
        self.assertIn("recorded", self.service.command("result 2.5 good patience"))
        self.assertIn("net +2.50", self.service.command(""))
        self.service.command("result -1")
        day = Plan30.load(self.store).day(1)
        self.assertEqual((day.net, day.manual), (-1.0, True))
        self.assertEqual(day.note, "good patience")   # a bare result keeps the note
    def test_result_validates_the_day_and_the_number(self):
        self.started()
        self.assertIn("/plan30", self.service.command("result"))
        self.assertIn("/plan30", self.service.command("result abc"))
        self.assertIn("not part of the plan", self.service.command("result 5 1.0"))

    def test_note_is_saved_for_today(self):
        self.started()
        self.assertIn("Note saved for day 1", self.service.command("note waited for the retest"))
        self.assertEqual(Plan30.load(self.store).day(1).note, "waited for the retest")
        self.assertIn("/plan30", self.service.command("note"))

    def test_table_lists_the_days_so_far(self):
        self.started()
        self.at(1)
        text = self.service.command("table")
        self.assertIn("D1 10-26", text)
        self.assertIn("D2 10-27", text)
        self.assertNotIn("D3", text)
        self.assertIn("/plan30 export", text)

    def test_unknown_subcommand_shows_usage(self):
        self.assertIn("/plan30 start", self.service.command("bogus"))

    # ----------------------------------------------------------- the day ends
    def test_a_finished_day_is_reviewed_once_and_then_preserved(self):
        self.started()
        self.store.record_fill(fill("a", profit=1.0, commission=-0.1))
        self.store.record_fill(fill("b", hour=11, profit=-0.4))
        self.at(1)
        self.service.step()
        review = self.messages()[-1]
        self.assertIn("Day 1 review", review)
        self.assertIn("100.00 → 100.50 USD (+0.50%)", review)
        self.assertIn("Trades 2 (won 1 / lost 1)", review)
        self.assertIn("all followed", review)
        self.assertIn("continue as planned", review)
        day = Plan30.load(self.store).day(1)
        self.assertTrue(day.final)
        count = len(self.messages())
        self.at(1, hour=13)
        self.service.step()
        self.assertEqual(len(self.messages()), count, "reviewed twice")
        self.assertIn("preserved", self.service.command("result 1 99"))
        self.assertEqual(Plan30.load(self.store).day(1).net, 0.5)

    def test_a_day_without_trades_is_a_no_trade_day_not_a_made_up_one(self):
        self.started()
        self.at(1)
        self.service.step()
        self.assertIn("No-trade day", self.messages()[-1])
        day = Plan30.load(self.store).day(1)
        self.assertEqual((day.net, day.trades), (0.0, 0))
        self.assertIn("No-trade day", day.note)

    def test_the_next_day_starts_from_where_the_last_really_ended(self):
        self.started()
        self.store.record_fill(fill(profit=-2.0))
        self.at(1)
        self.service.step()
        row = plan_module.rows(Plan30.load(self.store))[1]
        self.assertEqual(row.actual_start, 98.0)
        self.assertEqual(row.theoretical_start, 125.0)

    def test_a_broken_rule_is_named_and_risk_is_reduced(self):
        self.started()
        self.store.record_fill(fill(profit=-1.8))
        self.at(1)
        self.service.step()
        review = self.messages()[-1]
        self.assertIn("a trade lost more than the per-trade risk", review)
        self.assertIn("reduce risk", review)

    def test_a_day_that_cannot_be_worked_out_waits_for_you_and_blocks_the_next(self):
        self.started()
        self.broker.set_rules(replace(RULES, epic="GOLD", value_per_point=None))
        self.service._values.clear()
        self.store.record_fill(fill(profit=None))
        self.at(2)
        self.service.step()
        self.assertIn("Record it with /plan30 result 1", self.messages()[-1])
        self.assertFalse(Plan30.load(self.store).day(1).final)
        self.assertFalse(Plan30.load(self.store).day(2).final)
        count = len(self.messages())
        self.at(2, hour=13)
        self.service.step()
        self.assertEqual(len(self.messages()), count, "asked twice")
        self.assertIn("recorded", self.service.command("result 1 -0.80 missed a stop"))
        loaded = Plan30.load(self.store)
        self.assertTrue(loaded.day(1).final)
        self.assertEqual(loaded.day(1).net, -0.8)
        self.assertEqual(plan_module.rows(loaded)[1].actual_start, 99.2)

    def test_the_final_report_arrives_once_after_day_30(self):
        self.started()
        self.at(31)
        self.service.step()
        reports = [m for m in self.messages() if "30-Day Report" in m]
        self.assertEqual(len(reports), 1)
        self.assertIn("30/30 days recorded, complete", reports[0])
        self.assertIn("Rules followed on 30 of 30", reports[0])
        self.at(32)
        self.service.step()
        self.assertEqual(len([m for m in self.messages() if "30-Day Report" in m]), 1)

    # ----------------------------------------------------------- the loss limit
    def test_reaching_the_daily_loss_limit_warns_once(self):
        self.started()
        self.store.record_fill(fill("a", profit=-1.6))
        self.store.record_fill(fill("b", hour=11, profit=-1.5))
        self.service.step()
        warnings = [m for m in self.messages() if "daily loss limit is reached" in m]
        self.assertEqual(len(warnings), 1)
        self.assertIn("Stop trading for today", warnings[0])
        self.at(0, hour=13)
        self.service.step()
        self.assertEqual(len([m for m in self.messages()
                              if "daily loss limit is reached" in m]), 1)

    def test_a_small_loss_does_not_warn(self):
        self.started()
        self.store.record_fill(fill(profit=-1.0))
        self.service.step()
        self.assertFalse([m for m in self.messages() if "loss limit is reached" in m])

    def test_open_losses_count_toward_the_limit(self):
        from tests.helpers import position
        self.started()
        open_trade = replace(position(), upl=-3.5)
        self.broker.seed_position(open_trade)
        self.service.step()
        self.assertTrue([m for m in self.messages() if "loss limit is reached" in m])

    def test_nothing_here_places_or_changes_a_trade(self):
        self.started()
        self.store.record_fill(fill(profit=-5.0))
        before = list(self.broker.positions())
        self.at(1)
        self.service.step()
        self.service.command("export")
        self.assertEqual(self.broker.positions(), before)

    # ----------------------------------------------------------- export
    def test_export_sends_the_spreadsheet(self):
        self.started()
        with tempfile.TemporaryDirectory() as directory:
            self.supervisor.config.report.output_dir = directory
            reply = self.service.command("export")
        self.assertEqual(reply, "")
        self.assertEqual(len(self.notifier.documents), 1)
        self.assertTrue(self.notifier.documents[0][0].endswith(".xlsx"))

    def test_the_help_lists_the_command_and_it_is_registered(self):
        self.assertIn("/plan30", self.supervisor.help_text())
        seen = {}
        self.supervisor.notifier.register = lambda name, handler: seen.__setitem__(name, handler)
        self.supervisor._register_commands()
        self.assertIn("plan30", seen)

    def test_arabic_has_every_string(self):
        from tmbot.i18n import CATALOG
        for key, entry in CATALOG.items():
            if key.startswith("p30."):
                self.assertTrue(entry.get("ar"), key)


class ExcelTests(unittest.TestCase):
    def build(self):
        p = plan()
        for number, net in enumerate([10.0, -4.0, 0.0], start=1):
            log = p.day(number)
            log.net, log.final, log.trades = net, True, 1 if net else 0
            log.wins, log.losses = (1, 0) if net > 0 else (0, 1) if net < 0 else (0, 0)
            log.costs, log.max_dd = (-0.1 if net else 0.0), (1.0 if net else None)
            log.decision = "continue"
        p.day(3).note = "No-trade day"
        return p

    def test_the_workbook_has_the_four_sheets_and_the_required_columns(self):
        from openpyxl import load_workbook
        from tmbot.analysis import plan30_xlsx
        with tempfile.TemporaryDirectory() as directory:
            path = plan30_xlsx.export(self.build(), Path(directory) / "p.xlsx")
            book = load_workbook(path)
        self.assertEqual(book.sheetnames[0], "Plan")
        self.assertEqual(set(book.sheetnames), {"Plan", "Settings", "Details", "Summary"})
        headers = [c.value for c in book["Plan"][4]][:11]
        self.assertEqual(headers[0], "Day")
        self.assertEqual(headers[-1], "Notes and Lessons Learned")
        self.assertEqual(len(headers), 11)
        self.assertEqual(book["Plan"].max_row, 34)

    def test_calculations_are_formulas_and_results_are_values(self):
        from openpyxl import load_workbook
        from tmbot.analysis import plan30_xlsx
        with tempfile.TemporaryDirectory() as directory:
            sheet = load_workbook(plan30_xlsx.export(self.build(), Path(directory) / "p.xlsx"))["Plan"]
        self.assertEqual(sheet["D5"].value, "=Settings!$B$1")
        self.assertEqual(sheet["D6"].value, "=F5")
        self.assertTrue(str(sheet["E5"].value).startswith("=ROUND"))
        self.assertTrue(str(sheet["H5"].value).startswith("=IF"))
        self.assertEqual(sheet["G5"].value, 10.0)
        self.assertIsNone(sheet["G8"].value, "a day with no result must stay blank")

    @unittest.skipUnless(shutil.which("soffice") or shutil.which("libreoffice"),
                         "LibreOffice is needed to recalculate the formulas")
    def test_the_formulas_recalculate_to_the_same_numbers_as_the_code(self):
        from openpyxl import load_workbook
        from tmbot.analysis import plan30_xlsx
        p = self.build()
        office = shutil.which("soffice") or shutil.which("libreoffice")
        with tempfile.TemporaryDirectory() as directory:
            source = plan30_xlsx.export(p, Path(directory) / "p.xlsx")
            out = Path(directory) / "out"
            subprocess.run([office, "--headless", "--convert-to", "xlsx", "--outdir", str(out),
                            str(source)], check=True, capture_output=True, timeout=120)
            book = load_workbook(out / "p.xlsx", data_only=True)
        sheet = book["Plan"]
        for index, row in enumerate(plan_module.rows(p)):
            r = 5 + index
            self.assertAlmostEqual(sheet.cell(r, 4).value, row.theoretical_start, 2)
            self.assertAlmostEqual(sheet.cell(r, 5).value, row.theoretical_profit, 2)
            self.assertAlmostEqual(sheet.cell(r, 6).value, row.theoretical_end, 2)
            if row.actual_end is None:
                self.assertIn(sheet.cell(r, 8).value, (None, ""))
            else:
                self.assertAlmostEqual(sheet.cell(r, 3).value, row.actual_start, 2)
                self.assertAlmostEqual(sheet.cell(r, 8).value, row.actual_end, 2)
                self.assertEqual(sheet.cell(r, 9).value, "Yes" if row.achieved else "No")
        s, figures = book["Summary"], plan_module.summary(p)
        self.assertAlmostEqual(s["B3"].value, figures["initial"], 2)
        self.assertAlmostEqual(s["B4"].value, figures["final"], 2)
        self.assertAlmostEqual(s["B5"].value, figures["net"], 2)
        self.assertAlmostEqual(s["B6"].value * 100, figures["return_pct"], 2)
        self.assertAlmostEqual(s["B7"].value, figures["theoretical_final"], 2)
        self.assertAlmostEqual(s["B8"].value, figures["difference"], 2)
        self.assertEqual((s["B9"].value, s["B10"].value), (figures["wins"], figures["losses"]))
        self.assertAlmostEqual(s["B11"].value * 100, figures["win_rate"], 1)
        self.assertAlmostEqual(s["B12"].value, figures["max_drawdown"], 2)
        self.assertAlmostEqual(s["B13"].value, figures["costs"], 2)
        self.assertEqual(s["B14"].value, figures["days_reviewed"])
        self.assertEqual(s["B15"].value, figures["days_compliant"])


class StorageTests(unittest.TestCase):
    def test_the_plan_survives_a_round_trip_and_a_restart(self):
        _, _, store, _ = build_supervisor()
        p = plan(capital=250.0)
        p.day(1).net, p.day(1).final, p.day(1).broken = 3.0, True, ["trade_risk"]
        p.limit_alerted = [2]
        p.save(store)
        loaded = Plan30.load(store)
        self.assertEqual(loaded.to_dict(), p.to_dict())
        self.assertEqual(loaded.day(1).broken, ["trade_risk"])

    def test_garbage_in_the_store_does_not_crash(self):
        _, _, store, _ = build_supervisor()
        store.set(plan_module.KEY, "not json")
        self.assertIsNone(Plan30.load(store))


if __name__ == "__main__":
    unittest.main()
