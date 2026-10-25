"""Excel export of the 30-day plan, with live formulas.

Four sheets: ``Settings`` (the numbers every formula reads), ``Plan`` (the
eleven required columns), ``Details`` (wins, losses, costs, drawdown and rule
checks per day) and ``Summary`` (the final report).  Type a result into the
Actual Net column and everything after it recalculates.

Only recorded results are written as values; days with no result yet stay
blank, so a spreadsheet never shows an invented profit.
"""

from __future__ import annotations

from pathlib import Path
from typing import Callable, Optional

from .plan30 import Plan30, rows

FIRST = 5                      # first data row on the Plan and Details sheets
HEADERS = [
    "Day", "Date", "Actual Starting Balance", "Theoretical Starting Balance",
    "Theoretical Target Profit (25%)", "Theoretical Ending Balance",
    "Actual Net Profit/Loss", "Actual Ending Balance", "Target Achieved",
    "Number of Trades", "Notes and Lessons Learned",
]
DETAIL_HEADERS = [
    "Day", "Wins", "Losses", "Trading costs", "Max drawdown % (closed trades)",
    "Loss limit hit", "Rules followed", "Rules broken", "Decision", "P/L source",
]
RULE_TEXT = {
    "loss_limit": "daily loss limit exceeded",
    "trade_risk": "a trade lost more than the per-trade risk",
}


def export(
    plan: Plan30,
    path: str | Path,
    *,
    decision_text: Optional[Callable[[str], str]] = None,
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    decision_text = decision_text or (lambda key: key)
    last = FIRST + plan.days - 1
    book = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DCE6F1")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")
    currency = '#,##0.00'

    # ---------------------------------------------------------------- Settings
    settings = book.active
    settings.title = "Settings"
    for number, (label, value, fmt) in enumerate([
        ("Starting capital", plan.capital, currency),
        ("Theoretical daily growth target", plan.target, "0.00%"),
        ("Duration (days)", plan.days, "0"),
        ("Daily loss limit (share of the day's starting balance)", plan.loss_limit, "0.00%"),
        ("Risk per trade (share of the actual balance)", plan.risk, "0.00%"),
        ("Start date", plan.start_date, "yyyy-mm-dd"),
        ("Currency", plan.currency, "@"),
    ], start=1):
        settings.cell(row=number, column=1, value=label).font = bold
        cell = settings.cell(row=number, column=2, value=value)
        cell.number_format = fmt
        cell.fill = input_fill
    settings["A9"] = ("The 25% a day is a mathematical scenario, not a forecast or a "
                      "realistic expectation. Never force trades or add risk to reach it.")
    settings["A10"] = ("Yellow cells are inputs. Theoretical columns are projections only "
                       "and never replace the actual balance.")
    settings.column_dimensions["A"].width = 58
    settings.column_dimensions["B"].width = 16

    # ---------------------------------------------------------------- Plan
    sheet = book.create_sheet("Plan")
    sheet["A1"] = "30-Day Trading Growth Plan"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = ("Target Achieved = the day's actual net profit reached the daily growth "
                   "target (25%) of that day's actual starting balance. Blank = no result yet.")
    for column, title in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=FIRST - 1, column=column, value=title)
        cell.font, cell.fill, cell.alignment = bold, head_fill, wrap
    table = rows(plan)
    for index, row in enumerate(table):
        r = FIRST + index
        sheet.cell(row=r, column=1, value=row.number)
        sheet.cell(row=r, column=2, value=f"=Settings!$B$6+A{r}-1").number_format = "yyyy-mm-dd"
        sheet.cell(row=r, column=3, value=(
            "=Settings!$B$1" if index == 0 else f'=IF(H{r - 1}="","",H{r - 1})'))
        sheet.cell(row=r, column=4, value="=Settings!$B$1" if index == 0 else f"=F{r - 1}")
        sheet.cell(row=r, column=5, value=f"=ROUND(D{r}*Settings!$B$2,2)")
        sheet.cell(row=r, column=6, value=f"=ROUND(D{r}+E{r},2)")
        net = sheet.cell(row=r, column=7, value=row.net)
        net.fill = input_fill
        sheet.cell(row=r, column=8, value=f'=IF(OR(G{r}="",C{r}=""),"",ROUND(C{r}+G{r},2))')
        sheet.cell(row=r, column=9, value=(
            f'=IF(OR(G{r}="",C{r}=""),"",IF(G{r}>=ROUND(C{r}*Settings!$B$2,2),"Yes","No"))'))
        trades = sheet.cell(row=r, column=10, value=row.trades if row.net is not None else None)
        trades.fill = input_fill
        notes = sheet.cell(row=r, column=11, value=row.note or None)
        notes.fill, notes.alignment = input_fill, wrap
        for column in (3, 4, 5, 6, 7, 8):
            sheet.cell(row=r, column=column).number_format = currency
    for column, width in zip("ABCDEFGHIJK", (6, 12, 16, 16, 16, 16, 16, 16, 12, 12, 48)):
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[FIRST - 1].height = 48
    sheet.freeze_panes = f"C{FIRST}"

    # ---------------------------------------------------------------- Details
    details = book.create_sheet("Details")
    details["A1"] = "Daily detail and rule checks"
    details["A1"].font = Font(bold=True, size=14)
    details["A2"] = ("Max drawdown covers closed trades only. 'estimated' means some profit was "
                     "price move x lots x contract value, not money the broker booked.")
    for column, title in enumerate(DETAIL_HEADERS, start=1):
        cell = details.cell(row=FIRST - 1, column=column, value=title)
        cell.font, cell.fill, cell.alignment = bold, head_fill, wrap
    for index, row in enumerate(table):
        r, log = FIRST + index, row.log
        known = row.net is not None
        details.cell(row=r, column=1, value=row.number)
        details.cell(row=r, column=2, value=log.wins if known else None)
        details.cell(row=r, column=3, value=log.losses if known else None)
        details.cell(row=r, column=4, value=log.costs if known else None).number_format = currency
        details.cell(row=r, column=5, value=log.max_dd if known else None)
        details.cell(row=r, column=6, value=("Yes" if log.loss_limit_hit else "No") if known else None)
        details.cell(row=r, column=7, value=(("No" if log.broken else "Yes")
                                              if known and log.final else None))
        details.cell(row=r, column=8, value="; ".join(RULE_TEXT.get(k, k) for k in log.broken) or None)
        details.cell(row=r, column=9, value=decision_text(log.decision) if log.decision else None)
        details.cell(row=r, column=10, value=(
            ("manual" if log.manual else "estimated" if log.estimated else "broker")
            if known and log.trades else ("no trades" if known else None)))
    for column, width in zip("ABCDEFGHIJ", (6, 8, 8, 14, 22, 14, 14, 40, 26, 12)):
        details.column_dimensions[column].width = width
    details.row_dimensions[FIRST - 1].height = 48

    # ---------------------------------------------------------------- Summary
    final = book.create_sheet("Summary")
    final["A1"] = "Final 30-Day Report"
    final["A1"].font = Font(bold=True, size=14)
    plan_g, plan_f = f"Plan!G{FIRST}:G{last}", f"Plan!F{last}"
    lines = [
        ("Initial capital", "=Settings!B1", currency),
        ("Final actual balance", f"=Settings!B1+SUM({plan_g})", currency),
        ("Total actual net profit/loss", f"=SUM({plan_g})", currency),
        ("Actual percentage return", "=IF(Settings!B1=0,\"\",B5/Settings!B1)", "0.00%"),
        ("Theoretical projected final balance", f"={plan_f}", currency),
        ("Difference (theoretical - actual)", "=B7-B4", currency),
        ("Winning trades", f"=SUM(Details!B{FIRST}:B{last})", "0"),
        ("Losing trades", f"=SUM(Details!C{FIRST}:C{last})", "0"),
        ("Win rate", '=IF(B9+B10=0,"",B9/(B9+B10))', "0.0%"),
        ("Maximum drawdown % (closed trades)", f'=IF(COUNT(Details!E{FIRST}:E{last})=0,"",MAX(Details!E{FIRST}:E{last}))', "0.00"),
        ("Total trading costs", f"=SUM(Details!D{FIRST}:D{last})", currency),
        ("Days reviewed", f'=COUNTIF(Details!G{FIRST}:G{last},"Yes")+COUNTIF(Details!G{FIRST}:G{last},"No")', "0"),
        ("Days with every rule followed", f'=COUNTIF(Details!G{FIRST}:G{last},"Yes")', "0"),
        ("Days over the daily loss limit", f'=COUNTIF(Details!F{FIRST}:F{last},"Yes")', "0"),
    ]
    for offset, (label, formula, fmt) in enumerate(lines):
        r = 3 + offset
        final.cell(row=r, column=1, value=label).font = bold
        cell = final.cell(row=r, column=2, value=formula)
        cell.number_format = fmt
    final["A19"] = ("Results are the quality of the process, not a promise of any return. "
                    "Costs and drawdown are shown only where the data exists.")
    final.column_dimensions["A"].width = 42
    final.column_dimensions["B"].width = 18

    book.move_sheet("Plan", offset=-1)
    book.active = 0
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    return output
